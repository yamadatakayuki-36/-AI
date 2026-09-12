"""アジェンダ（議事録）の組み立てとファイル出力（テキスト／Word）。

実際の安全衛生委員会の議事録（■前回振り返り／■気づき事項／■テーマ「◯◯」
背景＋確認・共有／次回開催）の構成に合わせている。時間配分付きの進行表ではなく、
この「引き継ぎ＋今回のテーマ」という実務の議事録スタイルを採用している。

前回振り返り・気づき事項は自由記述（1行＝1項目）。行に「→」を含めると、
項目とその対応・フォローアップの2段構成として出力される。
例：「9/20 熱中症飲料終了　山田 → QRコードは各自保管、来年も使用」
"""

from __future__ import annotations

import re
from calendar import Calendar
from dataclasses import dataclass, field
from datetime import date
from io import BytesIO

from data_loader import load_company, load_members, load_theme_history, member_labels

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:  # python-docx 未インストール時のフォールバック
    Document = None
    Pt = None


@dataclass
class NoteLine:
    """前回振り返り・気づき事項の1項目（本文＋任意のフォローアップ）。"""

    text: str
    followup: str | None = None


@dataclass
class AgendaDocument:
    meeting_no: int
    year_month: str
    theme: str
    company_name: str
    committee_name: str
    meeting_place: str
    meeting_date: str
    datetime_label: str
    attendees: list[str]
    theme_points: list[str]
    reference_links: list[dict] = field(default_factory=list)
    action_owner: str = ""
    due_label: str = ""
    background: str = ""
    previous_review: list[str] = field(default_factory=list)
    awareness_items: list[str] = field(default_factory=list)
    next_meeting_label: str = ""


def third_friday(year_month: str) -> date:
    """指定した年月（YYYY-MM）の第3金曜日を返す。"""
    year, month = (int(x) for x in year_month.split("-"))
    fridays = [d for d in Calendar().itermonthdates(year, month) if d.month == month and d.weekday() == 4]
    return fridays[2] if len(fridays) >= 3 else fridays[-1]


def _next_year_month(year_month: str) -> str:
    year, month = (int(x) for x in year_month.split("-"))
    if month == 12:
        return f"{year + 1}-01"
    return f"{year}-{month + 1:02d}"


def format_datetime_label(d: date) -> str:
    weekdays = ["月", "火", "水", "木", "金", "土", "日"]
    return f"{d.year}年{d.month}月{d.day}日（{weekdays[d.weekday()]}）14:00〜15:00"


def format_next_meeting_label(d: date) -> str:
    weekdays = ["月", "火", "水", "木", "金", "土", "日"]
    return f"{d.year}年{d.month}月{d.day}日（{weekdays[d.weekday()]}）14:00〜"


def safe_filename_theme(theme: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\s]+', "_", theme.strip())
    return cleaned[:40] or "theme"


def parse_note_lines(text: str) -> list[NoteLine]:
    """複数行テキストを NoteLine のリストに変換する。

    各行に「→」が含まれていれば、その前後を項目本文とフォローアップに分ける。
    """
    result: list[NoteLine] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if "→" in line:
            head, _, tail = line.partition("→")
            result.append(NoteLine(head.strip(), tail.strip() or None))
        else:
            result.append(NoteLine(line))
    return result


_DEFAULT_LINKS = [
    {"title": "厚生労働省 職場のあんぜんサイト", "url": "https://anzeninfo.mhlw.go.jp/"},
    {"title": "中央労働災害防止協会（JISHA）", "url": "https://www.jisha.or.jp/"},
]

_DEFAULT_POINTS_TEMPLATE = [
    "先月の労働災害・ヒヤリハット報告の共有",
    "テーマに関する現状の課題整理",
    "具体的な対策案の検討",
    "次回までのアクションと担当決め",
]


def build_agenda(
    year_month: str,
    theme: str,
    points: list[str] | None = None,
    reason: str = "",
) -> AgendaDocument:
    """アジェンダ（議事録）を組み立てる。

    points（テーマ固有の討議ポイント）とreason（提案理由＝背景）が渡された場合は
    それを使い、無い場合（自分でテーマを入力した場合やOpenAI提案など）は
    汎用テンプレート・空欄にフォールバックする。
    """
    company = load_company()
    meeting_day = third_friday(year_month)
    next_meeting_day = third_friday(_next_year_month(year_month))
    meeting_no = len(load_theme_history()) + 1
    roster = member_labels(load_members())
    theme_points = list(points) if points else list(_DEFAULT_POINTS_TEMPLATE)
    return AgendaDocument(
        meeting_no=meeting_no,
        year_month=year_month,
        theme=theme,
        company_name=company["company_name"],
        committee_name="安全衛生委員会",
        meeting_place="会議室A",
        meeting_date=meeting_day.isoformat(),
        datetime_label=format_datetime_label(meeting_day),
        attendees=roster,
        theme_points=theme_points,
        reference_links=list(_DEFAULT_LINKS),
        action_owner=roster[0] if roster else "",
        due_label="次回委員会まで",
        background=reason,
        previous_review=[],
        awareness_items=[],
        next_meeting_label=format_next_meeting_label(next_meeting_day),
    )


def agenda_to_text(doc: AgendaDocument) -> str:
    lines = [
        f"第{doc.meeting_no}回 {doc.company_name} {doc.committee_name} 議事録",
        f"日時：{doc.datetime_label}",
        f"場所：{doc.meeting_place}",
        f"出席者：{'、'.join(doc.attendees)}",
    ]

    if doc.previous_review:
        lines += ["", "■前回振り返り"]
        for note in parse_note_lines("\n".join(doc.previous_review)):
            lines.append(f"・{note.text}")
            if note.followup:
                lines.append(f"　→{note.followup}")

    if doc.awareness_items:
        lines += ["", "■気づき事項"]
        for note in parse_note_lines("\n".join(doc.awareness_items)):
            lines.append(f"・{note.text}")
            if note.followup:
                lines.append(f"　→{note.followup}")

    lines += ["", f"■テーマ「{doc.theme}」"]
    if doc.background:
        lines.append(f"背景：{doc.background}")
    lines += ["", "確認・共有"]
    lines += [f"・{p}" for p in doc.theme_points]
    lines += ["", f"担当：{doc.action_owner}　期限：{doc.due_label}"]

    if doc.next_meeting_label:
        lines += ["", f"次回：{doc.next_meeting_label}"]

    if doc.reference_links:
        lines += ["", "【参考資料】"]
        lines += [f"　・{link['title']}：{link['url']}" for link in doc.reference_links]
    return "\n".join(lines)


def _add_note_paragraphs(document, lines: list[str]) -> None:
    for note in parse_note_lines("\n".join(lines)):
        document.add_paragraph(note.text, style="List Bullet")
        if note.followup:
            p = document.add_paragraph(f"→ {note.followup}")
            if Pt is not None:
                p.paragraph_format.left_indent = Pt(28)


def agenda_to_docx_bytes(doc: AgendaDocument) -> bytes:
    if Document is None:
        # python-docx 未インストール時はテキストを返す（拡張子はそのまま.docxとして保存される）
        return agenda_to_text(doc).encode("utf-8")

    document = Document()
    document.add_heading(f"第{doc.meeting_no}回 {doc.company_name} {doc.committee_name} 議事録", level=1)
    document.add_paragraph(f"日時：{doc.datetime_label}")
    document.add_paragraph(f"場所：{doc.meeting_place}")
    document.add_paragraph(f"出席者：{'、'.join(doc.attendees)}")

    if doc.previous_review:
        document.add_heading("前回振り返り", level=2)
        _add_note_paragraphs(document, doc.previous_review)

    if doc.awareness_items:
        document.add_heading("気づき事項", level=2)
        _add_note_paragraphs(document, doc.awareness_items)

    document.add_heading(f"テーマ「{doc.theme}」", level=2)
    if doc.background:
        p = document.add_paragraph()
        p.add_run("背景：").bold = True
        p.add_run(doc.background)

    document.add_heading("確認・共有", level=3)
    for point in doc.theme_points:
        document.add_paragraph(point, style="List Bullet")
    document.add_paragraph(f"担当：{doc.action_owner}　期限：{doc.due_label}")

    if doc.next_meeting_label:
        document.add_heading("次回開催", level=2)
        document.add_paragraph(doc.next_meeting_label)

    if doc.reference_links:
        document.add_heading("参考資料", level=2)
        for link in doc.reference_links:
            document.add_paragraph(f"{link['title']}：{link['url']}")

    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()
