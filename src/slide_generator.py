"""アジェンダ（議事録）からのスライド（.pptx）生成。"""

from __future__ import annotations

from io import BytesIO

from agenda_generator import AgendaDocument, agenda_to_text, parse_note_lines

try:
    from pptx import Presentation
except ImportError:  # python-pptx 未インストール時のフォールバック
    Presentation = None


def _fill_notes(text_frame, lines: list[str]) -> None:
    notes = parse_note_lines("\n".join(lines))
    first = True
    for note in notes:
        line = note.text if not note.followup else f"{note.text} → {note.followup}"
        if first:
            text_frame.text = line
            first = False
        else:
            p = text_frame.add_paragraph()
            p.text = line


def agenda_to_pptx_bytes(doc: AgendaDocument) -> bytes:
    if Presentation is None:
        # python-pptx 未インストール時はテキストを返す（拡張子はそのまま.pptxとして保存される）
        return agenda_to_text(doc).encode("utf-8")

    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = f"第{doc.meeting_no}回 {doc.committee_name}"
    title_slide.placeholders[1].text = f"{doc.company_name}\n{doc.datetime_label}"

    if doc.previous_review:
        review_slide = prs.slides.add_slide(prs.slide_layouts[1])
        review_slide.shapes.title.text = "前回振り返り"
        _fill_notes(review_slide.placeholders[1].text_frame, doc.previous_review)

    if doc.awareness_items:
        awareness_slide = prs.slides.add_slide(prs.slide_layouts[1])
        awareness_slide.shapes.title.text = "気づき事項"
        _fill_notes(awareness_slide.placeholders[1].text_frame, doc.awareness_items)

    theme_slide = prs.slides.add_slide(prs.slide_layouts[1])
    theme_slide.shapes.title.text = f"テーマ：{doc.theme}"
    body = theme_slide.placeholders[1].text_frame
    remaining_points = list(doc.theme_points)
    if doc.background:
        body.text = f"背景：{doc.background}"
    elif remaining_points:
        body.text = remaining_points.pop(0)
    for point in remaining_points:
        p = body.add_paragraph()
        p.text = point

    action_slide = prs.slides.add_slide(prs.slide_layouts[1])
    action_slide.shapes.title.text = "アクション・次回"
    action_body = action_slide.placeholders[1].text_frame
    action_body.text = f"担当：{doc.action_owner}"
    p = action_body.add_paragraph()
    p.text = f"期限：{doc.due_label}"
    if doc.next_meeting_label:
        p2 = action_body.add_paragraph()
        p2.text = f"次回：{doc.next_meeting_label}"

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()
